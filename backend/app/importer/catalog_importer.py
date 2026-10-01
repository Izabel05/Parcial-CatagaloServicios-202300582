from __future__ import annotations

import argparse
import hashlib
import json
import os
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl
try:
    import psycopg
    from psycopg.rows import dict_row
    from psycopg.types.json import Jsonb
except ModuleNotFoundError:  # Permite validar el Excel con --dry-run sin una base instalada.
    psycopg = None  # type: ignore[assignment]
    dict_row = None  # type: ignore[assignment]
    Jsonb = None  # type: ignore[assignment,misc]

EXPECTED_LEVEL_1 = 12
EXPECTED_LEVEL_2 = 46
SERVICE_START_ROW = 5
SERVICE_END_ROW = 101
SOURCE_SHEET = "Servicios Externos"


@dataclass(frozen=True)
class ResolvedCell:
    value: Any
    source_range: str | None
    is_merged: bool
    is_anchor: bool


@dataclass
class Observation:
    severity: str
    reason: str
    message: str
    source_range: str
    code: str | None = None
    source_row_start: int | None = None
    source_row_end: int | None = None
    details: dict[str, Any] | None = None


class MergedCellResolver:
    """Resuelve solo dentro del rango combinado que contiene la celda."""

    def __init__(self, worksheet: Any) -> None:
        self._cells: dict[tuple[int, int], tuple[str, Any, bool]] = {}
        for merged_range in worksheet.merged_cells.ranges:
            range_name = str(merged_range)
            anchor_value = worksheet.cell(merged_range.min_row, merged_range.min_col).value
            for row in range(merged_range.min_row, merged_range.max_row + 1):
                for column in range(merged_range.min_col, merged_range.max_col + 1):
                    self._cells[(row, column)] = (
                        range_name,
                        anchor_value,
                        row == merged_range.min_row and column == merged_range.min_col,
                    )

    def resolve(self, worksheet: Any, row: int, column: int) -> ResolvedCell:
        merged = self._cells.get((row, column))
        if merged is None:
            return ResolvedCell(worksheet.cell(row, column).value, None, False, True)
        range_name, anchor_value, is_anchor = merged
        return ResolvedCell(anchor_value, range_name, True, is_anchor)


def text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def numeric(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise ValueError(f"Valor numérico inválido: {value!r}") from error


def active_value(value: Any) -> bool | None:
    normalized = text(value)
    if normalized is None:
        return None
    if normalized == "S":
        return True
    if normalized == "N":
        return False
    raise ValueError(f"Valor ACTIVO no reconocido: {normalized}")


def has_value(value: Any) -> bool:
    return text(value) is not None


def row_has_values(worksheet: Any, row: int) -> bool:
    return any(has_value(worksheet.cell(row, column).value) for column in range(1, worksheet.max_column + 1))


def source_range(row: int) -> str:
    return f"A{row}:L{row}"


def parse_workbook(input_path: Path, sheet_name: str) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]], list[Observation], int]:
    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"No existe la hoja {sheet_name!r}. Hojas disponibles: {workbook.sheetnames}")
    worksheet = workbook[sheet_name]
    resolver = MergedCellResolver(worksheet)
    observations: list[Observation] = []
    levels: dict[str, dict[str, Any]] = {}
    services: list[dict[str, Any]] = []
    omitted_count = 0
    last_explicit_level1: tuple[str, int] | None = None

    def observe(**kwargs: Any) -> None:
        observations.append(Observation(**kwargs))

    for row in range(SERVICE_START_ROW, SERVICE_END_ROW + 1):
        raw_code = text(worksheet.cell(row, 3).value)
        raw_level1_code = text(worksheet.cell(row, 1).value)
        if raw_level1_code is not None:
            last_explicit_level1 = (raw_level1_code, row)
        resolved_code = resolver.resolve(worksheet, row, 3)

        if raw_code is None and resolved_code.value is not None and resolved_code.is_merged and not resolved_code.is_anchor:
            omitted_count += 1
            observe(
                severity="INFO",
                reason="MERGED_CONTINUATION",
                message="Fila omitida: pertenece a una celda combinada cuyo valor principal está en otra fila.",
                source_range=resolved_code.source_range or source_range(row),
                code=text(resolved_code.value),
                source_row_start=row,
                source_row_end=row,
                details={"anchor_rule": "solo se importa la fila ancla del rango combinado"},
            )
            continue

        if raw_code is None:
            if row_has_values(worksheet, row):
                omitted_count += 1
                observe(
                    severity="WARNING",
                    reason="NO_SERVICE_CODE",
                    message="Fila omitida: contiene datos, pero no tiene un COD.N2 explícito.",
                    source_range=source_range(row),
                    source_row_start=row,
                    source_row_end=row,
                )
            continue

        code = raw_code
        level1 = resolver.resolve(worksheet, row, 1)
        level1_name = resolver.resolve(worksheet, row, 2)
        level1_code = text(level1.value)
        level1_display_name = text(level1_name.value)
        inferred_level1 = False
        if level1_code is None:
            # Regla explícita para el bloque final: SE.12.3 está inmediatamente
            # después de SE.12.2, conserva el mismo bloque y no tiene un rango
            # combinado propio. Se registra la inferencia para que sea auditable.
            if last_explicit_level1 and last_explicit_level1[1] == row - 1:
                level1_code = last_explicit_level1[0]
                inferred_level1 = True
                observe(
                    severity="WARNING",
                    reason="LEVEL_1_BLOCK_CONTINUATION",
                    message=(
                        f"COD.N2 {code} se asoció a COD.N1 {level1_code} por continuidad inmediata "
                        "del bloque final; la inferencia queda registrada y no se propaga fuera de esa regla."
                    ),
                    source_range=source_range(row),
                    code=code,
                    source_row_start=row,
                    source_row_end=row,
                    details={"previous_explicit_level_1_row": row - 1, "rule": "immediate_previous_explicit_level_1"},
                )
                level1_display_name = levels.get(level1_code, {}).get("name")
            else:
                level1_display_name = None
        if level1_code is None:
            omitted_count += 1
            observe(
                severity="ERROR",
                reason="MISSING_LEVEL_1",
                message="Servicio omitido: no tiene un COD.N1 resoluble dentro de su rango combinado.",
                source_range=source_range(row),
                code=code,
                source_row_start=row,
                source_row_end=row,
            )
            continue

        if level1_code not in levels:
            levels[level1_code] = {
                "code": level1_code,
                "original_code": level1_code,
                "name": level1_display_name or "",
                "name_variants": [level1_display_name] if level1_display_name else [],
                "source_row": row,
                "source_range": f"A{row}:B{row}",
                "transformations": {
                    "canonical_name_rule": "first_explicit_occurrence",
                    "merged_cells_resolved": [value for value in (level1.source_range, level1_name.source_range) if value],
                    "code_normalization": "trim_only",
                },
                "review_status": "REVIEW" if level1_display_name is None else "OK",
            }
        elif level1_display_name and level1_display_name not in levels[level1_code]["name_variants"]:
            levels[level1_code]["name_variants"].append(level1_display_name)
            observe(
                severity="WARNING",
                reason="LEVEL_1_NAME_CONFLICT",
                message=(
                    f"COD.N1 {level1_code} conserva como nombre canónico el primer valor "
                    f"{levels[level1_code]['name']!r}; se observó también {level1_display_name!r}."
                ),
                source_range=f"A{row}:B{row}",
                code=level1_code,
                source_row_start=row,
                source_row_end=row,
                details={
                    "canonical_name": levels[level1_code]["name"],
                    "observed_name": level1_display_name,
                    "rule": "first_explicit_occurrence",
                },
            )

        resolved_fields = {
            "name": resolver.resolve(worksheet, row, 4),
            "active": resolver.resolve(worksheet, row, 5),
            "service_class": resolver.resolve(worksheet, row, 6),
            "criticality": resolver.resolve(worksheet, row, 7),
            "service_type": resolver.resolve(worksheet, row, 8),
            "description": resolver.resolve(worksheet, row, 9),
            "metric": resolver.resolve(worksheet, row, 10),
            "minimum": resolver.resolve(worksheet, row, 11),
            "maximum": resolver.resolve(worksheet, row, 12),
        }
        missing_review_fields = [
            field_name
            for field_name in ("active", "service_class", "criticality", "service_type", "metric")
            if text(resolved_fields[field_name].value) is None
        ]
        if text(resolved_fields["name"].value) is None:
            missing_review_fields.append("name")
        if missing_review_fields:
            observe(
                severity="WARNING",
                reason="INCOMPLETE_ATTRIBUTES",
                message="Servicio importado con atributos desconocidos y estado REVIEW.",
                source_range=source_range(row),
                code=code,
                source_row_start=row,
                source_row_end=row,
                details={"missing_fields": missing_review_fields},
            )

        merged_cells = sorted({cell.source_range for cell in (level1, level1_name, resolved_code, *resolved_fields.values()) if cell.source_range})
        transformations = {
            "merged_cells_resolved": merged_cells,
            "code_normalization": "trim_only",
            "original_code": code,
            "stored_code": code,
            "canonical_level_1_name_rule": "first_explicit_occurrence",
        }
        if inferred_level1:
            transformations["level_1_code_inferred"] = "immediate_previous_explicit_level_1"
        if level1_code == "SE.12":
            transformations["level_1_name_conflict"] = "canonical=Suministrar Analitica; variant=Mantener Tableros de Control"

        service = {
            "code": code,
            "original_code": code,
            "level_1_code": level1_code,
            "name": text(resolved_fields["name"].value) or "",
            "active": active_value(resolved_fields["active"].value),
            "service_class": text(resolved_fields["service_class"].value),
            "criticality": text(resolved_fields["criticality"].value),
            "service_type": text(resolved_fields["service_type"].value),
            "description": text(resolved_fields["description"].value),
            "metric": text(resolved_fields["metric"].value),
            "minimum": numeric(resolved_fields["minimum"].value),
            "maximum": numeric(resolved_fields["maximum"].value),
            "review_status": "REVIEW" if missing_review_fields else "OK",
            "source_sheet": sheet_name,
            "source_row": row,
            "source_range": source_range(row),
            "transformations": transformations,
        }
        services.append(service)

    option_rows = [row for row in range(SERVICE_END_ROW + 1, worksheet.max_row + 1) if row_has_values(worksheet, row)]
    if option_rows:
        omitted_count += len(option_rows)
        first_row, last_row = min(option_rows), max(option_rows)
        observe(
            severity="INFO",
            reason="OPTION_LIST",
            message="Filas omitidas: contienen listas de opciones, no registros del catálogo.",
            source_range=f"E{first_row}:H{last_row}",
            source_row_start=first_row,
            source_row_end=last_row,
            details={"rows_omitted": len(option_rows), "rule": "solo A5:L101 es rango de servicios"},
        )

    if len(levels) != EXPECTED_LEVEL_1:
        raise ValueError(f"Control fallido: se esperaban {EXPECTED_LEVEL_1} niveles 1 y se obtuvieron {len(levels)}")
    if len(services) != EXPECTED_LEVEL_2:
        raise ValueError(f"Control fallido: se esperaban {EXPECTED_LEVEL_2} servicios nivel 2 y se obtuvieron {len(services)}")
    return levels, services, observations, omitted_count


def repo_root() -> Path:
    return Path(__file__).resolve().parents[3]


def execute_schema(connection: psycopg.Connection) -> None:
    schema_path = repo_root() / "backend" / "db" / "00_schema.sql"
    for statement in schema_path.read_text(encoding="utf-8").split(";\n"):
        statement = statement.strip()
        if statement:
            connection.execute(statement)


def json_equal(left: Any, right: Any) -> bool:
    return json.dumps(left or {}, sort_keys=True, ensure_ascii=False) == json.dumps(right or {}, sort_keys=True, ensure_ascii=False)


def upsert_level(connection: psycopg.Connection, level: dict[str, Any], sheet_name: str) -> str:
    existing = connection.execute(
        "SELECT code, name, name_variants, source_sheet, source_row, source_range, transformations, review_status FROM service_level_1 WHERE code=%s FOR UPDATE",
        (level["code"],),
    ).fetchone()
    values = (
        level["original_code"], level["name"], Jsonb(level["name_variants"]), sheet_name,
        level["source_row"], level["source_range"], Jsonb(level["transformations"]), level["review_status"],
    )
    if existing is None:
        connection.execute(
            """INSERT INTO service_level_1
            (code, original_code, name, name_variants, source_sheet, source_row, source_range, transformations, review_status, sort_order)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
            (level["code"], *values, level["sort_order"]),
        )
        return "created"
    changed = (
        existing["name"] != level["name"]
        or not json_equal(existing["name_variants"], level["name_variants"])
        or existing["source_sheet"] != sheet_name
        or existing["source_row"] != level["source_row"]
        or existing["source_range"] != level["source_range"]
        or not json_equal(existing["transformations"], level["transformations"])
        or existing["review_status"] != level["review_status"]
    )
    if changed:
        connection.execute(
            """UPDATE service_level_1 SET original_code=%s,name=%s,name_variants=%s,source_sheet=%s,
            source_row=%s,source_range=%s,transformations=%s,review_status=%s WHERE code=%s""",
            (*values, level["code"]),
        )
        return "updated"
    return "unchanged"


def upsert_service(connection: psycopg.Connection, service: dict[str, Any]) -> str:
    existing = connection.execute(
        """SELECT code, original_code, level_1_code, name, active, service_class, criticality, service_type,
        description, metric, minimum, maximum, review_status, source_sheet, source_row, source_range, transformations
        FROM services WHERE code=%s FOR UPDATE""",
        (service["code"],),
    ).fetchone()
    values = (
        service["original_code"], service["level_1_code"], service["name"], service["active"],
        service["service_class"], service["criticality"], service["service_type"], service["description"],
        service["metric"], service["minimum"], service["maximum"], service["review_status"],
        service["source_sheet"], service["source_row"], service["source_range"], Jsonb(service["transformations"]),
    )
    if existing is None:
        connection.execute(
            """INSERT INTO services
            (code, original_code, level_1_code, name, active, service_class, criticality, service_type, description,
             metric, minimum, maximum, review_status, source_sheet, source_row, source_range, transformations)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
            (service["code"], *values),
        )
        return "created"
    current = (
        existing["original_code"], existing["level_1_code"], existing["name"], existing["active"],
        existing["service_class"], existing["criticality"], existing["service_type"], existing["description"],
        existing["metric"], existing["minimum"], existing["maximum"], existing["review_status"],
        existing["source_sheet"], existing["source_row"], existing["source_range"], existing["transformations"],
    )
    changed = any(left != right for left, right in zip(current[:-1], values[:-1])) or not json_equal(current[-1], service["transformations"])
    if changed:
        connection.execute(
            """UPDATE services SET original_code=%s,level_1_code=%s,name=%s,active=%s,service_class=%s,
            criticality=%s,service_type=%s,description=%s,metric=%s,minimum=%s,maximum=%s,review_status=%s,
            source_sheet=%s,source_row=%s,source_range=%s,transformations=%s,updated_at=CURRENT_TIMESTAMP WHERE code=%s""",
            (*values, service["code"]),
        )
        return "updated"
    return "unchanged"


def run_import(input_path: Path, sheet_name: str, database_url: str) -> dict[str, Any]:
    if psycopg is None:
        raise RuntimeError("Falta psycopg. Instala las dependencias con: pip install -r backend/requirements.txt")
    levels, services, observations, omitted_count = parse_workbook(input_path, sheet_name)
    source_sha256 = hashlib.sha256(input_path.read_bytes()).hexdigest()
    summary: dict[str, Any] = {
        "status": "COMPLETED", "source_file": str(input_path), "source_sha256": source_sha256,
        "sheet": sheet_name, "physical_rows": SERVICE_END_ROW - SERVICE_START_ROW + 1,
        "created": 0, "updated": 0, "unchanged": 0, "omitted": omitted_count,
        "observed": len(observations), "level_1": len(levels), "level_2": len(services),
        "level_1_created": 0, "level_1_updated": 0,
    }
    with psycopg.connect(database_url, row_factory=dict_row) as connection:
        with connection.transaction():
            execute_schema(connection)
            run_id = connection.execute(
                """INSERT INTO import_runs (status,source_file,source_sha256,sheet_name,physical_rows)
                VALUES ('RUNNING',%s,%s,%s,%s) RETURNING id""",
                (str(input_path), source_sha256, sheet_name, summary["physical_rows"]),
            ).fetchone()["id"]
            for index, level in enumerate(levels.values(), start=1):
                level["sort_order"] = index
                result = upsert_level(connection, level, sheet_name)
                if result == "created":
                    summary["level_1_created"] += 1
                elif result == "updated":
                    summary["level_1_updated"] += 1
            for service in services:
                result = upsert_service(connection, service)
                if result == "created":
                    summary["created"] += 1
                elif result == "updated":
                    summary["updated"] += 1
                else:
                    summary["unchanged"] += 1
            for observation in observations:
                connection.execute(
                    """INSERT INTO import_observations
                    (import_run_id,severity,reason,code,message,source_sheet,source_range,source_row_start,source_row_end,details)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (run_id, observation.severity, observation.reason, observation.code, observation.message,
                     sheet_name, observation.source_range, observation.source_row_start, observation.source_row_end,
                     Jsonb(observation.details or {})),
                )
            connection.execute(
                """INSERT INTO catalog_import_control
                (id,sheet_name,physical_rows,expected_level_1_count,expected_level_2_count,imported_level_1_count,imported_level_2_count,last_run_id)
                VALUES (1,%s,%s,%s,%s,%s,%s,%s)
                ON CONFLICT (id) DO UPDATE SET sheet_name=EXCLUDED.sheet_name,physical_rows=EXCLUDED.physical_rows,
                expected_level_1_count=EXCLUDED.expected_level_1_count,expected_level_2_count=EXCLUDED.expected_level_2_count,
                imported_level_1_count=EXCLUDED.imported_level_1_count,imported_level_2_count=EXCLUDED.imported_level_2_count,
                last_run_id=EXCLUDED.last_run_id,imported_at=CURRENT_TIMESTAMP""",
                (sheet_name, summary["physical_rows"], EXPECTED_LEVEL_1, EXPECTED_LEVEL_2, len(levels), len(services), run_id),
            )
            connection.execute(
                """UPDATE import_runs SET status='COMPLETED',finished_at=CURRENT_TIMESTAMP,created_count=%s,updated_count=%s,
                unchanged_count=%s,omitted_count=%s,observed_count=%s,level_1_count=%s,level_2_count=%s,summary=%s WHERE id=%s""",
                (summary["created"], summary["updated"], summary["unchanged"], summary["omitted"], summary["observed"],
                 summary["level_1"], summary["level_2"], Jsonb(summary), run_id),
            )
            summary["run_id"] = run_id
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description="Importa Servicios Externos desde Excel a PostgreSQL de forma idempotente.")
    parser.add_argument("--input", default=str(repo_root() / "data" / "CatalogoServicios.xlsx"))
    parser.add_argument("--sheet", default=SOURCE_SHEET)
    parser.add_argument("--database-url", default=os.getenv("DATABASE_URL"))
    parser.add_argument("--dry-run", action="store_true", help="Procesa el Excel y muestra controles sin escribir en PostgreSQL.")
    args = parser.parse_args()
    if args.dry_run:
        levels, services, observations, omitted = parse_workbook(Path(args.input).resolve(), args.sheet)
        print(json.dumps({
            "status": "DRY_RUN",
            "source_file": str(Path(args.input).resolve()),
            "sheet": args.sheet,
            "physical_rows": SERVICE_END_ROW - SERVICE_START_ROW + 1,
            "would_import_level_1": len(levels),
            "would_import_level_2": len(services),
            "omitted": omitted,
            "observed": len(observations),
        }, ensure_ascii=False, indent=2))
        return
    from backend.app.config.settings import settings

    summary = run_import(Path(args.input).resolve(), args.sheet, args.database_url or settings.database_url)
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
