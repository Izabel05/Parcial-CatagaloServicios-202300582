import json
import time
import unittest
from concurrent.futures import ThreadPoolExecutor

from tests.integration import test_scenarios as scenarios

INTEGRATION_AVAILABLE = scenarios.INTEGRATION_AVAILABLE
TEST_DATABASE_URL = scenarios.TEST_DATABASE_URL


@unittest.skipUnless(INTEGRATION_AVAILABLE and TEST_DATABASE_URL, 'Requiere PostgreSQL exclusivo de pruebas')
class CmdbIntegrationTest(scenarios.ScenarioIntegrationTest):
    @classmethod
    def setUpClass(cls):
        from urllib.parse import urlparse
        if not urlparse(TEST_DATABASE_URL).path.endswith('_test'):
            raise RuntimeError('La CMDB de pruebas requiere una base cuyo nombre termine en _test')
        super().setUpClass()
        from backend.app.database.migrate import migrate
        cls.migrations = migrate(TEST_DATABASE_URL)
        cls.assert_migration_repeat = migrate(TEST_DATABASE_URL)
        cls.admin_headers = {'Authorization': 'Bearer ' + cls.client.post('/api/v1/auth/login', json={
            'identifier': 'integration-admin', 'password': 'Admin-Test-12345'}).json()['access_token']}
        cls.reader_headers = {'Authorization': 'Bearer ' + cls.client.post('/api/v1/auth/login', json={
            'identifier': 'integration-consulta', 'password': 'Consulta-Test-12345'}).json()['access_token']}

    def post(self, path, payload, status=201):
        response = self.client.post('/api/v1/cmdb' + path, headers=self.admin_headers, json=payload)
        self.assertEqual(response.status_code, status, response.text)
        return response.json()

    def test_cmdb_01_generacion_5000_idempotente(self):
        import psycopg
        from backend.app.services.cmdb_generator import generate_devices
        started = time.perf_counter()
        first = generate_devices(582, 5000, self.integration_admin_id, TEST_DATABASE_URL)
        second = generate_devices(582, 5000, self.integration_admin_id, TEST_DATABASE_URL)
        self.assertEqual(first['creados'], 5000)
        self.assertEqual(second['creados'], 0)
        self.assertEqual(second['actualizados'], 5000)
        self.assertEqual(second['dispositivos_lote'], 5000)
        self.assertEqual(self.assert_migration_repeat, [])
        with psycopg.connect(TEST_DATABASE_URL) as connection:
            self.assertEqual(connection.execute('SELECT count(*),count(DISTINCT etiqueta) FROM cmdb_dispositivos').fetchone(), (5000, 5000))
            self.assertEqual(connection.execute("SELECT count(*) FROM cmdb_ci WHERE atributos_tecnicos ? 'ip'").fetchone()[0], 1666)
        response = self.client.get('/api/v1/cmdb/ci?page=2&page_size=31&es_dispositivo=true', headers=self.reader_headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['total'], 5000)
        self.assertEqual(len(response.json()['items']), 31)
        print(json.dumps({'cmdb_generar_y_repetir_5000_segundos': round(time.perf_counter() - started, 3), 'entorno': 'Docker Python 3.12/PostgreSQL 16'}))

    def test_cmdb_02_permisos_filtros_y_detalle(self):
        self.assertEqual(self.client.get('/api/v1/cmdb/ci').status_code, 401)
        denied = self.client.post('/api/v1/cmdb/catalogos/ubicaciones', headers=self.reader_headers,
                                  json={'codigo': 'PROHIBIDO', 'nombre': 'No permitido'})
        self.assertEqual(denied.status_code, 403)
        start = time.perf_counter()
        page = self.client.get('/api/v1/cmdb/ci?search=000123&page_size=10', headers=self.reader_headers).json()
        self.assertEqual(page['total'], 1)
        ci = page['items'][0]
        detail = self.client.get(f"/api/v1/cmdb/ci/{ci['id']}", headers=self.reader_headers).json()
        self.assertEqual(detail['clasificacion_completa'][0]['nombre'], 'Hardware')
        self.assertIsNotNone(detail['dispositivo'])
        self.assertGreater(len(detail['servicios']), 0)
        self.assertGreater(len(detail['historial']), 0)
        print(json.dumps({'cmdb_busqueda_y_detalle_segundos': round(time.perf_counter() - start, 4), 'dispositivos': 5000}))

    def test_cmdb_03_clasificaciones_y_ubicaciones(self):
        root = self.post('/catalogos/clasificaciones', {'codigo': 'TEST-ROOT', 'nombre': 'Clase de prueba'})
        child = self.post('/catalogos/clasificaciones', {'codigo': 'TEST-CHILD', 'nombre': 'Subclase de prueba', 'padre_id': root['id']})
        cycle = self.client.put(f"/api/v1/cmdb/catalogos/clasificaciones/{root['id']}", headers=self.admin_headers,
                                json={'codigo': root['codigo'], 'nombre': root['nombre'], 'padre_id': child['id']})
        self.assertEqual(cycle.status_code, 409)
        self.post('/catalogos/clasificaciones', {'codigo': 'BAD', 'nombre': 'Inválido', 'padre_id': 9999999}, 404)
        self.post('/catalogos/clasificaciones', {'codigo': root['codigo'], 'nombre': 'Duplicado'}, 409)
        self.assertEqual(self.client.delete(f"/api/v1/cmdb/catalogos/clasificaciones/{root['id']}", headers=self.admin_headers).status_code, 409)
        self.assertEqual(self.client.delete(f"/api/v1/cmdb/catalogos/clasificaciones/{child['id']}", headers=self.admin_headers).status_code, 200)
        location = self.post('/catalogos/ubicaciones', {'codigo': 'EDIT', 'nombre': 'Antes'})
        changed = self.client.put(f"/api/v1/cmdb/catalogos/ubicaciones/{location['id']}", headers=self.admin_headers,
                                  json={'codigo': 'EDIT', 'nombre': 'Después'})
        self.assertEqual(changed.json()['nombre'], 'Después')

    def test_cmdb_04_relaciones_ciclos_e_impacto(self):
        rows = self.client.get('/api/v1/cmdb/ci?page_size=3', headers=self.admin_headers).json()['items']
        a, b, c = [r['id'] for r in rows]
        relation = self.post('/relaciones', {'origen_ci_id': a, 'destino_ci_id': b, 'tipo': 'depende_de', 'justificacion': 'Prueba de dependencia'})
        self.post('/relaciones', {'origen_ci_id': b, 'destino_ci_id': a, 'tipo': 'alojado_en', 'justificacion': 'Ciclo mixto'}, 409)
        self.post('/relaciones', {'origen_ci_id': a, 'destino_ci_id': a, 'tipo': 'conectado_a', 'justificacion': 'Autorreferencia'}, 409)
        for x, y in [(a, b), (b, c), (c, a)]:
            self.post('/relaciones', {'origen_ci_id': x, 'destino_ci_id': y, 'tipo': 'conectado_a', 'justificacion': 'Conexión cíclica permitida'})
        self.post('/relaciones', {'origen_ci_id': b, 'destino_ci_id': a, 'tipo': 'conectado_a', 'justificacion': 'Duplicado inverso'}, 409)
        impact = self.client.get(f'/api/v1/cmdb/ci/{b}/dependencias', headers=self.reader_headers).json()
        self.assertEqual(impact['alcanzados'], 2)
        self.assertTrue(any(s['ruta_ci'] == [b, a] for s in impact['servicios_potencialmente_afectados']))
        limited = self.client.get(f'/api/v1/cmdb/ci/{b}/dependencias?limit=1', headers=self.reader_headers).json()
        self.assertTrue(limited['truncado'])
        self.assertEqual(self.client.delete(f"/api/v1/cmdb/relaciones/{relation['id']}", headers=self.reader_headers).status_code, 403)
        edited = self.client.put(f"/api/v1/cmdb/relaciones/{relation['id']}", headers=self.admin_headers,
                                 json={'origen_ci_id': a, 'destino_ci_id': c, 'tipo': 'depende_de', 'justificacion': 'Edición trazable'})
        self.assertEqual(edited.status_code, 200, edited.text)
        self.assertEqual(self.client.delete(f"/api/v1/cmdb/relaciones/{relation['id']}", headers=self.admin_headers).status_code, 200)

    def test_cmdb_05_mantenimiento_ci_y_referencias(self):
        source = self.client.get('/api/v1/cmdb/ci?page_size=1', headers=self.admin_headers).json()['items'][0]
        payload = {'codigo': 'MANUAL-CI', 'nombre': 'Equipo manual', 'clasificacion_id': source['clasificacion_id'],
                   'ubicacion_id': source['ubicacion_id'], 'responsable_usuario_id': self.integration_admin_id,
                   'etiqueta': 'MANUAL-TAG', 'atributos_tecnicos': {'memoria_gb': 16},
                   'relaciones_originales': 'Base de Datos1', 'documentacion_relacionada': '[Manual](url)',
                   'fecha_cambio_original': '2022-02-01'}
        ci = self.post('/ci', payload)
        self.assertEqual(ci['relaciones_originales'], 'Base de Datos1')
        self.assertEqual(ci['fecha_cambio_original'], '2022-02-01')
        self.post('/ci', {**payload, 'codigo': 'OTRO-CI'}, 409)
        bad = self.client.post('/api/v1/cmdb/ci', headers=self.admin_headers,
                               json={**payload, 'atributos_tecnicos': {'ip': 'no-ip'}})
        self.assertEqual(bad.status_code, 422)
        self.post('/ci', {**payload, 'codigo': 'BAD-REF', 'etiqueta': 'BAD-REF', 'ubicacion_id': 999999}, 404)
        service_id = self.client.get('/api/v1/services?page_size=1', headers=self.admin_headers).json()['items'][0]['id']
        self.post(f"/ci/{ci['id']}/servicios", {'servicio_id': service_id}, 200)
        self.assertEqual(self.client.delete(f"/api/v1/cmdb/ci/{ci['id']}", headers=self.admin_headers).status_code, 409)
        self.assertEqual(self.client.delete(f"/api/v1/cmdb/ci/{ci['id']}/servicios/{service_id}", headers=self.admin_headers).status_code, 200)
        edited = self.client.put(f"/api/v1/cmdb/ci/{ci['id']}", headers=self.admin_headers, json={**payload, 'nombre': 'Editado'})
        self.assertEqual(edited.status_code, 200, edited.text)
        self.assertEqual(self.client.delete(f"/api/v1/cmdb/ci/{ci['id']}", headers=self.admin_headers).status_code, 200)
        history = self.client.get(f"/api/v1/cmdb/ci/{ci['id']}", headers=self.reader_headers).json()['historial']
        self.assertTrue(all(h['actor_id'] == self.integration_admin_id for h in history))
        self.assertIn('editar', {h['accion'] for h in history})

    def test_cmdb_06_ciclos_concurrentes(self):
        from backend.app.services.cmdb_service import CmdbService, CmdbError
        from backend.app.repositories.cmdb_repository import CmdbRepository
        a = self.post('/catalogos/clasificaciones', {'codigo': 'CONC-A', 'nombre': 'Concurrente A'})
        b = self.post('/catalogos/clasificaciones', {'codigo': 'CONC-B', 'nombre': 'Concurrente B'})
        def move(pair):
            x, y = pair
            try:
                CmdbService(CmdbRepository(database_url=TEST_DATABASE_URL)).save_catalog('clasificaciones',
                    {'codigo': x['codigo'], 'nombre': x['nombre'], 'padre_id': y['id'], 'activo': True},
                    self.integration_admin_id, x['id'])
                return True
            except CmdbError:
                return False
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(move, [(a, b), (b, a)]))
        self.assertEqual(sum(results), 1)

    def test_cmdb_07_correspondencia_fuente_disponible(self):
        import hashlib
        from pathlib import Path
        from openpyxl import load_workbook
        source=Path('data/fuentes/CMDB.xlsx')
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(),'d182796bb40b492be2cb87e2c5fe19a8f910f2e754183dc48b5e89bd360464dc')
        workbook=load_workbook(source,read_only=True,data_only=False,keep_links=False)
        try:
            self.assertEqual(workbook.sheetnames,['CMDB'])
            sheet=workbook['CMDB']
            self.assertEqual(len([sheet.cell(4,c).value for c in range(2,22)]),20)
            labels={str(sheet['C5'].value).strip(),str(sheet['C6'].value).strip()}
            self.assertEqual(labels,{'Hardware','Software'})
        finally:workbook.close()
        rows=self.client.get('/api/v1/cmdb/catalogos/clasificaciones',headers=self.admin_headers).json()
        for label,cell in [('Hardware','CMDB!C5'),('Software','CMDB!C6')]:
            row=next(r for r in rows if r['codigo']=='FUENTE-'+label.upper())
            self.assertEqual(row['nombre'],label);self.assertEqual(row['referencia_fuente'],cell)
            self.assertEqual(row['origen'],'fuente');self.assertIsNone(row['padre_id'])
        print(json.dumps({'P04':'Correspondencia con C5/C6 y atributos B4:U4 verificada; la fuente no contiene árbol multinivel',
                          'jerarquia_multinivel_verificada':False},ensure_ascii=False))

    def test_cmdb_08_mediciones_filtros_5000(self):
        rows=self.client.get('/api/v1/cmdb/catalogos/clasificaciones',headers=self.admin_headers).json()
        classification=next(r['id'] for r in rows if r['codigo']=='FUENTE-HARDWARE')
        samples=[]
        for query,expected_size in [('page=2&page_size=25&es_dispositivo=true&activo=true',25),
                ('search=000123&page_size=10&activo=true',1),
                (f'clasificacion_id={classification}&es_dispositivo=true&activo=true&page_size=17',17)]:
            start=time.perf_counter()
            response=self.client.get('/api/v1/cmdb/ci?'+query,headers=self.reader_headers)
            elapsed=time.perf_counter()-start
            self.assertEqual(response.status_code,200,response.text)
            self.assertEqual(len(response.json()['items']),expected_size)
            samples.append({'consulta':query,'segundos':round(elapsed,6),'filas':len(response.json()['items']),'total':response.json()['total']})
        print(json.dumps({'P16':samples,'entorno':'FastAPI TestClient, Docker Python 3.12/PostgreSQL 16, 5000 dispositivos'},ensure_ascii=False))


if __name__ == '__main__':
    unittest.main()
